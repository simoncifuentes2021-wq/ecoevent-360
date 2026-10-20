from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from uuid import UUID

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.permissions import can_manage_event, can_view_operational_event
from app.db.session import set_document_event_context
from app.models.documents import DocumentEmployment, DocumentPerson, DocumentSignatureRequest, ManagedDocument
from app.models.core import Event, User
from app.models.enums import UserRole
from app.schemas.document_schema import (
    EmploymentInput,
    EmploymentUpdate,
    ImportPreview,
    PersonCreate,
    PersonImportCommit,
    PersonImportRow,
    PersonUpdate,
)

IMPORT_COLUMNS = (
    "RUT",
    "NOMBRE",
    "APELLIDO",
    "EMAIL",
    "TELEFONO",
    "DOMICILIO",
    "COMUNA",
    "CARGO",
    "AREA",
    "LUGAR_TRABAJO",
    "FECHA_INICIO",
    "FECHA_TERMINO",
    "REMUNERACION",
    "LABORES_FUNCIONES",
)
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 500


def normalize_rut(value: str) -> str:
    raw = str(value or "").strip().upper()
    match = re.fullmatch(r"(\d{1,8}|\d{1,3}(?:\.\d{3})+)-([0-9K])", raw)
    if not match:
        raise ValueError("RUT inválido")
    # Leading zeroes are mathematically neutral in the verifier calculation,
    # but canonicalizing them keeps display_rut() round-trippable.
    digits = match.group(1).replace(".", "").lstrip("0") or "0"
    verifier = match.group(2)
    if len(digits) > 8:
        raise ValueError("RUT inválido")
    total, multiplier = 0, 2
    for digit in reversed(digits):
        total += int(digit) * multiplier
        multiplier = 2 if multiplier == 7 else multiplier + 1
    result = 11 - total % 11
    expected = "0" if result == 11 else "K" if result == 10 else str(result)
    if verifier != expected:
        raise ValueError("RUT inválido")
    return f"{digits}{verifier}"


_LOWERCASE_NAME_PARTICLES = {"de", "del", "la", "las", "los", "y", "e"}


def normalize_person_name(value: str) -> str:
    """Trim repeated whitespace and apply readable Spanish name capitalization."""
    normalized = unicodedata.normalize("NFC", " ".join(str(value or "").split())).lower()
    words = normalized.split(" ")
    formatted: list[str] = []
    for index, word in enumerate(words):
        if index > 0 and word in _LOWERCASE_NAME_PARTICLES:
            formatted.append(word)
            continue
        characters: list[str] = []
        capitalize_next = True
        for character in word:
            if character in {"-", "'"}:
                characters.append(character)
                capitalize_next = True
            else:
                characters.append(character.upper() if capitalize_next else character)
                capitalize_next = False
        formatted.append("".join(characters))
    return " ".join(formatted)


def display_rut(normalized: str) -> str:
    return f"{int(normalized[:-1]):,}".replace(",", ".") + f"-{normalized[-1]}"


def resolve_client_id(
    db: Session, user: User, requested: UUID | None, event_id: UUID | None = None
) -> UUID:
    event = db.get(Event, event_id) if event_id else None
    if event_id and not event:
        raise HTTPException(status_code=404, detail="Evento no encontrado")
    if event and not can_view_operational_event(user, event, db):
        raise HTTPException(status_code=403, detail="No tienes acceso a este evento")
    client_id = requested or (event.client_id if event else user.client_id)
    if client_id is None:
        raise HTTPException(
            status_code=422, detail="Selecciona el cliente al que pertenece la persona"
        )
    if event and event.client_id != client_id:
        raise HTTPException(status_code=422, detail="El evento no pertenece al cliente indicado")
    if user.role not in {
        UserRole.SUPER_ADMIN,
        UserRole.ADMIN,
        UserRole.SUPERVISOR,
    }:
        raise HTTPException(status_code=403, detail="Rol sin acceso al módulo Documentos")
    return client_id


def _validate_date_range(start: date | None, end: date | None) -> None:
    if start and end and end < start:
        raise HTTPException(
            status_code=422, detail="La fecha de término debe ser posterior al inicio"
        )


def create_person(db: Session, payload: PersonCreate, user: User) -> DocumentPerson:
    employment = payload.employment
    if not employment or not employment.event_id:
        raise HTTPException(
            status_code=422, detail="Asocia a la persona a una faena/evento y define su cargo"
        )
    client_id = resolve_client_id(
        db, user, payload.client_id, employment.event_id if employment else None
    )
    try:
        normalized = normalize_rut(payload.rut)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="El RUT ingresado no es válido") from exc
    if user.role == UserRole.SUPERVISOR and not can_manage_event(user, employment.event_id, db):
        raise HTTPException(status_code=403, detail="No puedes administrar personas en este evento")
    if user.role == UserRole.SUPERVISOR:
        set_document_event_context(db, employment.event_id)
        db.connection().execute(
            text("select set_config('app.document_person_lookup_rut', :rut, true)"),
            {"rut": normalized},
        )
    _validate_date_range(employment.start_date, employment.end_date)
    person = db.scalar(
        select(DocumentPerson)
        .where(DocumentPerson.rut_normalized == normalized)
        .with_for_update()
    )
    if person is None:
        person = DocumentPerson(
            first_name=normalize_person_name(payload.first_name),
            last_name=normalize_person_name(payload.last_name),
            rut=display_rut(normalized),
            rut_normalized=normalized,
            email=str(payload.email).lower() if payload.email else None,
            phone=payload.phone.strip() if payload.phone else None,
            address=payload.address.strip() if payload.address else None,
            commune=payload.commune.strip() if payload.commune else None,
        )
        db.add(person)
        db.flush()
    else:
        person.first_name = normalize_person_name(payload.first_name)
        person.last_name = normalize_person_name(payload.last_name)
        person.rut = display_rut(normalized)
        person.email = str(payload.email).lower() if payload.email else person.email
        person.phone = payload.phone.strip() if payload.phone else person.phone
        person.address = payload.address.strip() if payload.address else person.address
        person.commune = payload.commune.strip() if payload.commune else person.commune
        person.updated_at = datetime.utcnow()
    existing_assignment = db.scalar(
        select(DocumentEmployment.id).where(
            DocumentEmployment.person_id == person.id,
            DocumentEmployment.event_id == employment.event_id,
            DocumentEmployment.position == employment.position.strip(),
            DocumentEmployment.start_date == employment.start_date,
        )
    )
    if existing_assignment:
        raise HTTPException(status_code=409, detail="La persona ya está asignada a esta faena con esos datos")
    db.add(_employment(person, client_id, employment))
    db.commit()
    db.refresh(person)
    return person


def _person_for_update(
    db: Session, person_id: UUID, user: User, event_id: UUID | None
) -> DocumentPerson:
    if user.role not in {UserRole.SUPER_ADMIN, UserRole.ADMIN}:
        raise HTTPException(status_code=403, detail="Solo admin puede editar la ficha reutilizable")

    person = db.scalar(
        select(DocumentPerson).where(DocumentPerson.id == person_id).with_for_update()
    )
    if not person:
        raise HTTPException(status_code=404, detail="Persona no encontrada")
    if event_id:
        assignment_id = db.scalar(
            select(DocumentEmployment.id).where(
                DocumentEmployment.person_id == person_id,
                DocumentEmployment.event_id == event_id,
            ).limit(1)
        )
        if not assignment_id:
            raise HTTPException(status_code=404, detail="La persona no pertenece a esa faena")
    return person


def update_person(
    db: Session,
    person_id: UUID,
    payload: PersonUpdate,
    user: User,
    *,
    event_id: UUID | None = None,
) -> DocumentPerson:
    person = _person_for_update(db, person_id, user, event_id)
    try:
        normalized = normalize_rut(payload.rut)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="El RUT ingresado no es vÃ¡lido") from exc
    duplicate = db.scalar(
        select(DocumentPerson.id).where(
            DocumentPerson.rut_normalized == normalized,
            DocumentPerson.id != person.id,
        ).limit(1)
    )
    if duplicate:
        raise HTTPException(status_code=409, detail="Ese RUT ya pertenece a otra ficha")
    if normalized != person.rut_normalized:
        active_link = db.scalar(
            select(DocumentSignatureRequest.id)
            .join(ManagedDocument, ManagedDocument.id == DocumentSignatureRequest.document_id)
            .where(
                ManagedDocument.person_id == person.id,
                DocumentSignatureRequest.revoked_at.is_(None),
                DocumentSignatureRequest.expires_at > datetime.utcnow(),
            )
            .limit(1)
        )
        if active_link:
            raise HTTPException(
                status_code=409,
                detail="No se puede cambiar el RUT mientras haya enlaces de firma vigentes. Reenvía o revoca esos enlaces primero.",
            )

    person.first_name = normalize_person_name(payload.first_name)
    person.last_name = normalize_person_name(payload.last_name)
    person.rut = display_rut(normalized)
    person.rut_normalized = normalized
    person.email = str(payload.email).lower() if payload.email else None
    person.phone = payload.phone.strip() if payload.phone else None
    person.address = payload.address.strip() if payload.address else None
    person.commune = payload.commune.strip() if payload.commune else None
    person.updated_at = datetime.utcnow()
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="No se pudo guardar: revisa si el RUT ya existe") from exc
    db.refresh(person)
    return person


def update_employment(
    db: Session,
    person_id: UUID,
    employment_id: UUID,
    payload: EmploymentUpdate,
    user: User,
    *,
    event_id: UUID | None = None,
) -> DocumentEmployment:
    if user.role not in {UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.SUPERVISOR}:
        raise HTTPException(status_code=403, detail="Rol sin acceso al mÃ³dulo Documentos")
    if user.role == UserRole.SUPERVISOR:
        if not event_id or not can_manage_event(user, event_id, db):
            raise HTTPException(status_code=403, detail="Selecciona una faena que tengas asignada")
        set_document_event_context(db, event_id)

    assignment = db.scalar(
        select(DocumentEmployment)
        .where(
            DocumentEmployment.id == employment_id,
            DocumentEmployment.person_id == person_id,
        )
        .with_for_update()
    )
    if not assignment or (event_id and assignment.event_id != event_id):
        raise HTTPException(status_code=404, detail="AsignaciÃ³n de faena no encontrada")
    if user.role == UserRole.SUPERVISOR and assignment.event_id != event_id:
        raise HTTPException(status_code=403, detail="No puedes editar otra faena")
    _validate_date_range(payload.start_date, payload.end_date)

    assignment.position = payload.position.strip()
    assignment.area = payload.area.strip() if payload.area else None
    assignment.start_date = payload.start_date
    assignment.end_date = payload.end_date
    assignment.remuneration = payload.remuneration
    assignment.workplace = payload.workplace.strip() if payload.workplace else None
    assignment.duties = payload.duties.strip() if payload.duties else None
    assignment.updated_at = datetime.utcnow()
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=422, detail="Revisa las fechas de la asignaciÃ³n") from exc
    db.refresh(assignment)
    return assignment


def _employment(
    person: DocumentPerson,
    client_id: UUID,
    data: EmploymentInput | PersonImportRow,
    event_id: UUID | None = None,
) -> DocumentEmployment:
    return DocumentEmployment(
        client_id=client_id,
        person_id=person.id,
        event_id=event_id or getattr(data, "event_id", None),
        position=data.position.strip(),
        area=data.area.strip() if data.area else None,
        start_date=data.start_date,
        end_date=data.end_date,
        remuneration=data.remuneration,
        workplace=getattr(data, "workplace", None),
        duties=getattr(data, "duties", None),
    )


def list_people(
    db: Session,
    user: User,
    *,
    q: str | None = None,
    event_id: UUID | None = None,
    client_id: UUID | None = None,
) -> list[dict]:
    stmt = (
        select(DocumentPerson)
        .order_by(DocumentPerson.last_name, DocumentPerson.first_name)
        .limit(500)
    )
    if user.role in {UserRole.SUPER_ADMIN, UserRole.ADMIN}:
        if client_id:
            stmt = stmt.join(DocumentEmployment).where(DocumentEmployment.client_id == client_id)
    elif user.role == UserRole.SUPERVISOR:
        if not event_id or not can_manage_event(user, event_id, db):
            raise HTTPException(status_code=403, detail="Selecciona un evento que tengas asignado")
        stmt = stmt.join(DocumentEmployment).where(DocumentEmployment.event_id == event_id)
    else:
        raise HTTPException(status_code=403, detail="Rol sin acceso al módulo Documentos")
    if event_id and user.role not in {UserRole.SUPERVISOR}:
        stmt = stmt.join(DocumentEmployment).where(DocumentEmployment.event_id == event_id)
    if q:
        needle = f"%{q.strip()}%"
        stmt = stmt.where(
            (DocumentPerson.first_name.ilike(needle))
            | (DocumentPerson.last_name.ilike(needle))
            | (DocumentPerson.rut.ilike(needle))
            | (DocumentPerson.email.ilike(needle))
        )
    people = db.scalars(stmt.distinct()).all()
    result = []
    for person in people:
        assignment_query = select(DocumentEmployment).where(DocumentEmployment.person_id == person.id)
        if event_id:
            assignment_query = assignment_query.where(DocumentEmployment.event_id == event_id)
        elif client_id:
            assignment_query = assignment_query.where(DocumentEmployment.client_id == client_id)
        assignment = db.scalar(
            assignment_query.order_by(DocumentEmployment.created_at.desc())
            .limit(1)
        )
        result.append(
            {
                "id": person.id,
                "first_name": person.first_name,
                "last_name": person.last_name,
                "rut": person.rut,
                "email": person.email,
                "phone": person.phone,
                "address": person.address,
                "commune": person.commune,
                "updated_at": person.updated_at,
                "latest_employment": (
                    {
                        "id": assignment.id,
                        "event_id": assignment.event_id,
                        "position": assignment.position,
                        "area": assignment.area,
                        "workplace": assignment.workplace,
                        "duties": assignment.duties,
                        "start_date": assignment.start_date,
                        "end_date": assignment.end_date,
                        "remuneration": str(assignment.remuneration)
                        if assignment.remuneration is not None
                        else None,
                    }
                    if assignment
                    else None
                ),
            }
        )
    return result


def _header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    normalized = (
        "".join(ch for ch in text if not unicodedata.combining(ch))
        .strip()
        .upper()
        .replace(" ", "_")
    )
    if normalized in {"DEPENDENCIAS", "DEPENDENCIAS_DE_TRABAJO", "LUGAR_DE_TRABAJO"}:
        return "LUGAR_TRABAJO"
    if normalized == "FAENA":
        return "LUGAR_TRABAJO"
    if normalized in {"LABORES", "FUNCIONES", "LABORES_Y_FUNCIONES", "LABORES_FUNCION"}:
        return "LABORES_FUNCIONES"
    return normalized


def _parse_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    raise ValueError("Fecha inválida; usa DD/MM/AAAA")


def preview_import(content: bytes) -> ImportPreview:
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="El Excel supera el máximo de 5 MB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [_header(v) for v in next(rows, ())]
    except Exception as exc:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo Excel") from exc
    indexes = {name: headers.index(name) for name in IMPORT_COLUMNS if name in headers}
    missing = [name for name in IMPORT_COLUMNS[:3] + IMPORT_COLUMNS[7:8] if name not in indexes]
    if missing:
        workbook.close()
        raise HTTPException(
            status_code=422, detail=f"Faltan columnas obligatorias: {', '.join(missing)}"
        )
    valid: list[dict] = []
    errors: list[dict] = []
    seen: set[str] = set()
    total = 0
    for row_number, cells in enumerate(rows, start=2):
        if row_number > MAX_IMPORT_ROWS + 1:
            workbook.close()
            raise HTTPException(status_code=422, detail="El Excel no puede superar 500 filas")
        values = {
            field: (
                cells[indexes[field]] if field in indexes and indexes[field] < len(cells) else None
            )
            for field in IMPORT_COLUMNS
        }
        if not any(value not in (None, "") for value in values.values()):
            continue
        total += 1
        row_errors: list[dict] = []
        try:
            rut_norm = normalize_rut(str(values["RUT"] or ""))
        except ValueError:
            rut_norm = ""
            row_errors.append({"row": row_number, "field": "RUT", "message": "RUT inválido"})
        if rut_norm and rut_norm in seen:
            row_errors.append(
                {"row": row_number, "field": "RUT", "message": "RUT duplicado dentro del Excel"}
            )
        if rut_norm:
            seen.add(rut_norm)
        try:
            raw_amount = values["REMUNERACION"]
            amount = (
                (
                    Decimal(str(raw_amount))
                    if isinstance(raw_amount, (int, float, Decimal))
                    else Decimal(
                        str(raw_amount).replace("$", "").replace(".", "").replace(",", ".")
                    )
                )
                if raw_amount not in (None, "")
                else None
            )
            start = _parse_date(values["FECHA_INICIO"])
            end = _parse_date(values["FECHA_TERMINO"])
            if start and end and end < start:
                raise ValueError("La fecha de término debe ser posterior al inicio")
            person = PersonImportRow(
                first_name=normalize_person_name(str(values["NOMBRE"] or "")),
                last_name=normalize_person_name(str(values["APELLIDO"] or "")),
                rut=display_rut(rut_norm) if rut_norm else str(values["RUT"] or "").strip(),
                email=str(values["EMAIL"]).strip() if values["EMAIL"] else None,
                phone=str(values["TELEFONO"]).strip() if values["TELEFONO"] else None,
                address=str(values["DOMICILIO"]).strip() if values["DOMICILIO"] else None,
                commune=str(values["COMUNA"]).strip() if values["COMUNA"] else None,
                position=str(values["CARGO"] or "").strip(),
                area=str(values["AREA"]).strip() if values["AREA"] else None,
                workplace=(
                    str(values["LUGAR_TRABAJO"]).strip()
                    if values["LUGAR_TRABAJO"]
                    else None
                ),
                start_date=start,
                end_date=end,
                remuneration=amount,
                duties=(
                    " ".join(str(values["LABORES_FUNCIONES"]).split()).strip()
                    if values["LABORES_FUNCIONES"] not in (None, "")
                    else None
                ),
            )
            normalize_rut(person.rut)
        except (ValueError, InvalidOperation, ValidationError) as exc:
            text = str(exc)
            field = "FECHA" if "fecha" in text.lower() else "DATOS"
            row_errors.append({"row": row_number, "field": field, "message": text[:180]})
            person = None
        if row_errors:
            errors.extend(row_errors)
        elif person:
            valid.append(person.model_dump(mode="json"))
    workbook.close()
    return ImportPreview(
        valid_rows=valid,
        errors=errors,
        total_rows=total,
        valid_count=len(valid),
        error_count=total - len(valid),
    )


def commit_import(db: Session, payload: PersonImportCommit, user: User) -> dict:
    if not payload.event_id:
        raise HTTPException(status_code=422, detail="Selecciona la faena/evento al que corresponde la importación")
    client_id = resolve_client_id(db, user, payload.client_id, payload.event_id)
    if user.role == UserRole.SUPERVISOR and (
        not payload.event_id or not can_manage_event(user, payload.event_id, db)
    ):
        raise HTTPException(
            status_code=403, detail="El supervisor debe importar al evento asignado"
        )
    if user.role == UserRole.SUPERVISOR and payload.event_id:
        set_document_event_context(db, payload.event_id)
    normalized_rows: list[tuple[PersonImportRow, str]] = []
    seen: set[str] = set()
    for row in payload.rows:
        try:
            normalized = normalize_rut(row.rut)
        except ValueError as exc:
            raise HTTPException(
                status_code=422, detail="Uno de los RUT de importación no es válido"
            ) from exc
        if normalized in seen:
            raise HTTPException(
                status_code=409,
                detail="El Excel contiene RUT duplicados; corrige la vista previa antes de importar",
            )
        seen.add(normalized)
        _validate_date_range(row.start_date, row.end_date)
        normalized_rows.append((row, normalized))
    created = updated = assignments = 0
    results = []
    try:
        for row, normalized in normalized_rows:
            try:
                with db.begin_nested():
                    person = db.scalar(
                        select(DocumentPerson)
                        .where(DocumentPerson.rut_normalized == normalized)
                        .with_for_update()
                    )
                    if person:
                        person.first_name, person.last_name, person.rut = (
                            normalize_person_name(row.first_name),
                            normalize_person_name(row.last_name),
                            display_rut(normalized),
                        )
                        person.email = str(row.email).lower() if row.email else None
                        person.phone, person.address, person.commune = row.phone, row.address, row.commune
                        person.updated_at = datetime.utcnow()
                        person_action = "updated"
                    else:
                        person = DocumentPerson(
                            first_name=normalize_person_name(row.first_name),
                            last_name=normalize_person_name(row.last_name),
                            rut=display_rut(normalized),
                            rut_normalized=normalized,
                            email=str(row.email).lower() if row.email else None,
                            phone=row.phone,
                            address=row.address,
                            commune=row.commune,
                        )
                        db.add(person)
                        db.flush()
                        person_action = "created"
                    assignment = db.scalar(
                        select(DocumentEmployment)
                        .where(
                            DocumentEmployment.client_id == client_id,
                            DocumentEmployment.person_id == person.id,
                            DocumentEmployment.event_id == payload.event_id
                            if payload.event_id
                            else DocumentEmployment.event_id.is_(None),
                            DocumentEmployment.position == row.position.strip(),
                            DocumentEmployment.start_date == row.start_date,
                        )
                        .order_by(DocumentEmployment.created_at.desc())
                        .limit(1)
                        .with_for_update()
                    )
                    if assignment:
                        assignment.area = row.area.strip() if row.area else None
                        assignment.workplace = row.workplace.strip() if row.workplace else assignment.workplace
                        assignment.end_date = row.end_date
                        assignment.remuneration = row.remuneration
                        assignment.duties = row.duties.strip() if row.duties and row.duties.strip() else assignment.duties
                        assignment_action = "updated"
                    else:
                        db.add(_employment(person, client_id, row, payload.event_id))
                        assignment_action = "created"
                    db.flush()
                created += person_action == "created"
                updated += person_action == "updated"
                assignments += 1
                results.append({
                    "person_id": str(person.id),
                    "rut": display_rut(normalized),
                    "name": f"{person.first_name} {person.last_name}",
                    "person_action": person_action,
                    "assignment_action": assignment_action,
                    "status": "SUCCESS",
                    "error": None,
                })
            except Exception:
                results.append({
                    "rut": display_rut(normalized),
                    "name": f"{normalize_person_name(row.first_name)} {normalize_person_name(row.last_name)}",
                    "person_action": "unchanged",
                    "assignment_action": "unchanged",
                    "status": "ERROR",
                    "error": "No se pudo procesar esta fila. Puedes reintentarla individualmente.",
                })
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {
        "created": created,
        "updated": updated,
        "assignments": assignments,
        "workers_processed": len(normalized_rows),
        "workers_failed": sum(row["status"] == "ERROR" for row in results),
        "results": results,
    }


def excel_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Personas"
    template_columns = tuple(
        "FAENA" if column == "LUGAR_TRABAJO" else column for column in IMPORT_COLUMNS
    )
    sheet.append(template_columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F5B")
    sheet["A1"].comment = Comment(
        "RUT con guion obligatorio; puntos opcionales. Ejemplos: 11111111-1 o 11.111.111-1.",
        "Greenway App",
    )
    sheet["N1"].comment = Comment(
        "Opcional. Se incluirá en la cláusula PRIMERO del contrato como labores/funciones. "
        "Si se deja vacío, el contrato usará el Cargo como hasta ahora.",
        "Greenway App",
    )
    for column, width in enumerate(
        (16, 20, 22, 28, 18, 32, 20, 22, 20, 34, 18, 18, 18, 42), start=1
    ):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(IMPORT_COLUMNS))}1"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()

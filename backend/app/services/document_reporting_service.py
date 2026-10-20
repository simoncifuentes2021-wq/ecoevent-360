from io import BytesIO
from uuid import UUID

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.permissions import can_view_operational_event
from app.models.core import Event, User
from app.services import document_people_service, managed_documents_service

_STATUS_LABELS = {
    "DRAFT": "Borrador",
    "READY": "Listo para enviar",
    "SENT": "Enviado",
    "VIEWED": "Visto",
    "PARTIALLY_SIGNED": "Firma parcial",
    "COMPLETED": "Completado",
    "CANCELLED": "Cancelado",
    "EXPIRED": "Vencido",
    "SUPERSEDED": "Reemplazado",
}


def build_event_document_status_xlsx(db: Session, event_id: UUID, user: User) -> bytes:
    event = db.get(Event, event_id)
    if not event or not can_view_operational_event(user, event, db):
        raise HTTPException(status_code=404, detail="Faena no encontrada o sin acceso")

    people = document_people_service.list_people(db, user, event_id=event_id)
    documents = managed_documents_service.list_documents(db, user, event_id=event_id)
    latest_by_person: dict[tuple[UUID, str], dict] = {}
    for document in documents:
        key = (document["person_id"], document["kind"])
        previous = latest_by_person.get(key)
        if previous is None or (
            document["version"], document["updated_at"]
        ) > (previous["version"], previous["updated_at"]):
            latest_by_person[key] = document

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Estado documental"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:M1")
    sheet["A1"] = "GREENWAY APP  |  ESTADO DOCUMENTAL"
    sheet["A1"].font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="087F68")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells("A2:M2")
    sheet["A2"] = f"Faena: {event.name}   ·   Personas: {len(people)}"
    sheet["A2"].font = Font(size=11, bold=True, color="173B35")
    sheet["A2"].alignment = Alignment(vertical="center")

    completed_contracts = sum(
        latest_by_person.get((person["id"], "CONTRACT"), {}).get("status") == "COMPLETED"
        for person in people
    )
    completed_epp = sum(
        latest_by_person.get((person["id"], "EPP"), {}).get("status") == "COMPLETED"
        for person in people
    )
    metrics = [
        ("PERSONAS", len(people)),
        ("CONTRATOS COMPLETADOS", completed_contracts),
        ("EPP COMPLETADOS", completed_epp),
        ("DOCUMENTOS FALTANTES", sum(
            (person["id"], kind) not in latest_by_person
            for person in people for kind in ("CONTRACT", "EPP")
        )),
    ]
    for index, (label, value) in enumerate(metrics):
        column = index * 3 + 1
        label_cell = sheet.cell(row=4, column=column, value=label)
        value_cell = sheet.cell(row=5, column=column, value=value)
        label_cell.font = Font(size=9, bold=True, color="53656F")
        value_cell.font = Font(size=15, bold=True, color="087F68")
        value_cell.alignment = Alignment(vertical="center")

    headers = [
        "Faena", "Nombre completo", "RUT", "Correo electrónico", "Teléfono",
        "Cargo", "Área", "Contrato · estado", "Contrato · versión",
        "Registro EPP · estado", "EPP · versión", "Fecha de inicio", "Fecha de término",
    ]
    header_row = 7
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="145A4B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 30

    for row_number, person in enumerate(people, start=header_row + 1):
        assignment = person.get("latest_employment") or {}
        contract = latest_by_person.get((person["id"], "CONTRACT"))
        epp = latest_by_person.get((person["id"], "EPP"))
        values = [
            event.name,
            f"{person['first_name']} {person['last_name']}",
            person["rut"],
            person.get("email") or "—",
            person.get("phone") or "—",
            assignment.get("position") or "—",
            assignment.get("area") or "—",
            _STATUS_LABELS.get(contract["status"], contract["status"]) if contract else "Faltante",
            f"v{contract['version']}" if contract else "—",
            _STATUS_LABELS.get(epp["status"], epp["status"]) if epp else "Faltante",
            f"v{epp['version']}" if epp else "—",
            assignment.get("start_date") or "—",
            assignment.get("end_date") or "—",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if column in {8, 10}:
                cell.font = Font(
                    bold=True,
                    color="A16207" if value == "Faltante" else "087F68" if value == "Completado" else "334155",
                )
        sheet.row_dimensions[row_number].height = 30

    widths = [28, 30, 17, 32, 17, 24, 22, 23, 16, 25, 14, 17, 17]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    last_row = header_row + len(people)
    sheet.freeze_panes = "A8"
    sheet.auto_filter.ref = f"A{header_row}:M{last_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.print_title_rows = f"1:{header_row}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()

from io import BytesIO
import hashlib
import json

from openpyxl import Workbook
import pytest
from fastapi import HTTPException
from pypdf import PdfReader

from app.services.document_people_service import (
    IMPORT_COLUMNS,
    display_rut,
    normalize_person_name,
    normalize_rut,
    preview_import,
)
from app.core.config import settings
from app.models.enums import UserRole
from app.services.document_people_service import resolve_client_id
from app.services import managed_documents_service
from app.services.document_storage_service import (
    ConfiguredPrivateDocumentStorage,
    LocalPrivateDocumentStorage,
)
from app.services.report_pdf_service import render_html
from app.services.managed_documents_service import (
    _apply_worker_epp_acknowledgments,
    _copy_epp_delivery,
    _epp_data,
    _read,
    render_document_html,
    token_digest,
    preview_pdf,
)
from app.services.signature_provider import LocalSimulatedSignatureProvider
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from uuid import uuid4
from types import SimpleNamespace


def _workbook(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(IMPORT_COLUMNS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_chilean_rut_is_normalized_and_displayed_without_changing_identity():
    assert normalize_rut("11.111.111-1") == "111111111"
    assert normalize_rut("11111111-1") == "111111111"
    assert display_rut("111111111") == "11.111.111-1"
    assert normalize_rut("09.361.610-3") == normalize_rut("9.361.610-3")
    assert normalize_rut(display_rut(normalize_rut("09.361.610-3"))) == normalize_rut("09.361.610-3")


@pytest.mark.parametrize(
    "rut",
    [
        "111111111",
        "11.111.1111-1",
        "11x.111.111-1",
        "11..111.111-1",
        "11.111.111/1",
    ],
)
def test_chilean_rut_requires_hyphen_and_rejects_malformed_characters(rut):
    with pytest.raises(ValueError, match="RUT inválido"):
        normalize_rut(rut)


def test_person_names_are_trimmed_and_formatted_for_documents():
    assert normalize_person_name("  sIMON   ") == "Simon"
    assert normalize_person_name("cIFUENTES pÉREZ") == "Cifuentes Pérez"
    assert normalize_person_name("juan-pABLO") == "Juan-Pablo"


def test_document_portal_links_resolve_to_configured_or_production_frontend(monkeypatch):
    monkeypatch.setattr(settings, "documents_public_base_url", None)
    monkeypatch.setattr(settings, "public_app_url", "https://cliente.example/")
    monkeypatch.setattr(settings, "app_env", "production")
    assert settings.resolved_documents_public_base_url == "https://cliente.example"

    monkeypatch.setattr(settings, "public_app_url", None)
    assert settings.resolved_documents_public_base_url == "https://app.greenway.cl"

    monkeypatch.setattr(settings, "app_env", "local")
    assert settings.resolved_documents_public_base_url == "http://localhost:3001"


def test_epp_draft_without_frozen_snapshot_can_be_listed():
    document = SimpleNamespace(
        epp_delivery=SimpleNamespace(
            riohs_received=False,
            eirl_completed=False,
            delivery_date=None,
            supervisor_id=None,
            items=[],
        ),
        frozen_snapshot=None,
    )

    assert _epp_data(document) == {
        "riohs_received": False,
        "eirl_completed": False,
        "delivery_date": None,
        "supervisor_id": None,
        "items": [],
    }


@pytest.mark.parametrize("kind", ["CONTRACT", "EPP"])
def test_admin_preview_includes_worker_signature_before_administrative_closure(kind, monkeypatch):
    signature_data_url = "data:image/png;base64,d29ya2VyLXNpZ25hdHVyZQ=="
    document = SimpleNamespace(
        id=uuid4(),
        signed_pdf_key=None,
        frozen_snapshot={"kind": kind},
        kind=kind,
        signatures=[SimpleNamespace(
            revoked_at=None,
            worker_signed_at=datetime(2026, 10, 16, 12, 0),
            worker_signed_name="Camila Rojas",
            worker_form_data={"worker_signature_data_url": signature_data_url},
            created_at=datetime(2026, 10, 16, 12, 0),
        )],
    )

    class Query:
        def where(self, *_args):
            return self

    class Session:
        def scalar(self, _query):
            return document

    monkeypatch.setattr(managed_documents_service, "_document_query", Query)
    monkeypatch.setattr(managed_documents_service, "_allowed", lambda *_args: None)
    monkeypatch.setattr(managed_documents_service, "render_document_html", lambda _snapshot, signatures: json.dumps(signatures))
    monkeypatch.setattr(managed_documents_service, "render_html", lambda html: (html.encode(), None))
    monkeypatch.setattr(
        managed_documents_service,
        "_apply_worker_epp_acknowledgments",
        lambda snapshot, _ack: snapshot,
    )

    preview = preview_pdf(Session(), document.id, SimpleNamespace())
    signatures = json.loads(preview)
    assert signatures["worker_name"] == "Camila Rojas"
    assert signatures["worker_signed_at"] == "2026-10-16T12:00:00"
    assert signatures["worker_signature_data_url"] == signature_data_url


@pytest.mark.parametrize("role", [UserRole.ADMIN, UserRole.SUPER_ADMIN])
def test_admin_roles_can_manage_documents_across_clients(role):
    selected_client = uuid4()
    admin = SimpleNamespace(role=role, client_id=None)
    assert resolve_client_id(None, admin, selected_client) == selected_client


def test_admin_must_select_client_when_creating_global_person():
    admin = SimpleNamespace(role=UserRole.ADMIN, client_id=None)
    with pytest.raises(HTTPException) as error:
        resolve_client_id(None, admin, None)
    assert error.value.status_code == 422


def test_excel_preview_accepts_valid_rows_and_rejects_duplicate_rut_and_invalid_email():
    valid = [
        "11111111-1",
        "sIMON",
        "cIFUENTES",
        "ana@example.cl",
        "+56912345678",
        "Calle 1",
        "Santiago",
        "Operaria",
        "Recinto",
        "Bodega central",
        "01/10/2026",
        "31/10/2026",
        850000,
    ]
    duplicate = [
        "11.111.111-1",
        "Otra",
        "Persona",
        "otro@example.cl",
        "",
        "",
        "",
        "Operario",
        "",
        "",
        "",
        "",
        850000,
    ]
    invalid_email = [
        "12.345.678-5",
        "Mal",
        "Correo",
        "esto-no-es-email",
        "",
        "",
        "",
        "Técnico",
        "",
        "",
        "",
        "",
        0,
    ]
    result = preview_import(_workbook([valid, duplicate, invalid_email]))
    assert result.total_rows == 3
    assert result.valid_count == 1
    assert result.error_count == 2
    assert any(error.field == "RUT" and error.row == 3 for error in result.errors)
    assert any(error.row == 4 for error in result.errors)
    assert result.valid_rows[0]["remuneration"] == "850000"
    assert result.valid_rows[0]["workplace"] == "Bodega central"
    assert result.valid_rows[0]["rut"] == "11.111.111-1"
    assert result.valid_rows[0]["first_name"] == "Simon"
    assert result.valid_rows[0]["last_name"] == "Cifuentes"


def test_excel_preview_reads_optional_contract_duties_and_keeps_blank_optional():
    with_duties = [
        "11.111.111-1", "Ana", "Pérez", "ana@example.cl", "", "", "", "Staff", "", "",
        "01/10/2026", "02/10/2026", 85000, "  Staff de aseo,   limpieza y reciclaje  ",
    ]
    without_duties = [
        "12.345.678-5", "Luis", "Pérez", "luis@example.cl", "", "", "", "Operario", "", "",
        "01/10/2026", "02/10/2026", 85000, None,
    ]
    result = preview_import(_workbook([with_duties, without_duties]))
    assert result.valid_count == 2
    assert result.valid_rows[0]["duties"] == "Staff de aseo, limpieza y reciclaje"
    assert result.valid_rows[1]["duties"] is None


def test_excel_template_includes_optional_duties_column():
    from openpyxl import load_workbook

    from app.services.document_people_service import excel_template

    workbook = load_workbook(BytesIO(excel_template()))
    sheet = workbook.active
    assert sheet["J1"].value == "FAENA"
    assert sheet["N1"].value == "LABORES_FUNCIONES"
    assert "Si se deja vacío" in sheet["N1"].comment.text
    assert sheet.auto_filter.ref == "A1:N1"


def test_excel_preview_reports_invalid_dates_by_row_and_ignores_blank_rows():
    invalid_date = [
        "11.111.111-1",
        "Ana",
        "Pérez",
        "ana@example.cl",
        "",
        "",
        "",
        "Operaria",
        "",
        "",
        "ayer",
        "",
        "",
    ]
    result = preview_import(_workbook([invalid_date, [None] * len(IMPORT_COLUMNS)]))
    assert result.total_rows == 1
    assert result.valid_count == 0
    assert result.errors[0].row == 2
    assert "Fecha" in result.errors[0].message


def _contract_snapshot(*, event=None):
    snapshot = {
        "title": "Contrato de prueba",
        "version": 1,
        "kind": "CONTRACT",
        "prepared_at": "2026-10-06",
        "person": {
            "first_name": "Ana",
            "last_name": "Pérez",
            "rut": "11.111.111-1",
            "address": "Calle Uno 123",
            "commune": "Santiago",
        },
        "employment": {
            "position": "Staff de aseo, limpieza y reciclaje",
            "start_date": "2026-09-27",
            "end_date": "2026-10-03",
            "remuneration": "85000",
            "workplace": "Recinto Principal",
        },
        "event": event,
        "data": {
            "contract_date": "2026-09-26",
            "employer_name": "Servicios Prueba SPA",
            "employer_rut": "76.123.456-7",
            "representative_name": "Persona Representante",
            "representative_rut": "12.345.678-5",
            "employer_address": "Avenida Uno 123",
            "employer_commune": "Santiago",
        },
        "signature_notice": "Firma simulada local",
    }
    return snapshot


def test_contract_uses_source_clauses_and_dynamic_fields_without_event():
    snapshot = _contract_snapshot()
    snapshot["data"]["employer_name"] = "Servicios <script>alert(1)</script>"
    html = render_document_html(snapshot)
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "<script>alert(1)</script>" not in html
    assert "CONTRATO POR OBRA O FAENA" in html
    assert "En Santiago, a <strong>26 de Septiembre de 2026</strong>" in html
    assert "PRIMERO:" in html and "Staff de aseo, limpieza y reciclaje" in html
    assert "SEGUNDO:" in html and "85.000</strong> diarios" in html
    assert "TERCERO:" in html and "CUARTO:" in html and "QUINTO:" in html
    assert "SEXTO:" in html and "SEPTIMO:" in html and "OCTAVO:" in html and "NOVENO:" in html
    assert "el evento denominado <strong>________________________</strong>" in html
    assert "en las dependencias de <strong>Recinto Principal</strong>" in html
    assert "No constituye por sí mismo" not in html


def test_contract_event_name_and_location_are_rendered_from_event_snapshot():
    html = render_document_html(
        _contract_snapshot(event={"name": "Festival Local", "venue": "Parque Central"})
    )
    assert "evento denominado <strong>Festival Local</strong>" in html
    assert "dependencias de <strong>Parque Central</strong>" in html


def test_paper_document_templates_keep_blank_employer_signature_space():
    contract = _contract_snapshot()
    contract_html = render_document_html(contract)
    assert "Firma manuscrita del trabajador" in contract_html
    assert "Firma del empleador / representante legal" in contract_html
    assert "class='signature-image'" not in contract_html

    epp_html = render_document_html(
        {
            "kind": "EPP",
            "person": {"first_name": "Ana", "last_name": "Pérez"},
            "epp": {"items": []},
        }
    )
    assert "Firma del Trabajador" in epp_html
    assert "Firma / Timbre del Instructor o Supervisor" in epp_html
    assert "class='signature-image'" not in epp_html


def test_existing_lowercase_names_are_formatted_in_contract_and_epp_documents():
    contract = _contract_snapshot()
    contract["person"]["first_name"] = "sIMON"
    contract["person"]["last_name"] = "cIFUENTES pÉREZ"
    contract_html = render_document_html(contract)
    assert "don (a) <strong>Simon Cifuentes Pérez</strong>" in contract_html

    epp = {
        "kind": "EPP",
        "person": {"first_name": "sIMON", "last_name": "cIFUENTES pÉREZ"},
        "employment": {},
        "data": {},
        "epp": {"items": []},
    }
    epp_html = render_document_html(epp)
    assert "NOMBRE TRABAJADOR : <strong>Simon Cifuentes Pérez</strong>" in epp_html


def test_epp_template_keeps_source_text_catalog_logo_and_dynamic_company_data():
    snapshot = {
        "title": "Registro EPP",
        "kind": "EPP",
        "person": {"first_name": "Ana", "last_name": "Pérez", "rut": "11.111.111-1"},
        "employment": {"position": "Staff", "area": "Limpieza"},
        "employer_name": "Servicios Prueba SPA",
        "data": {},
        "epp": {"items": [{"label": "Guantes", "received_at": "2026-10-01", "received_conformity": True}]},
        "signature_notice": "Firma simulada local",
    }
    html = render_document_html(snapshot)
    assert 'src="data:image/png;base64,' in html
    assert "REGISTRO DE RECEPCIÓN DE DOCUMENTOS Y ELEMEMENTOS DE PROTECCIÓN PERSONAL" in html
    assert "ELEMENENTOS DE PROTECCIÓN PERSONAL" in html
    assert "ELEMENTO ENTREGADO" in html and "FECHA DE RECEPCIÓN" in html
    assert "RECIBÍ CONFORME" in html and "FECHA DE DEVOLUCIÓN" not in html
    assert "Guantes" in html and "Jockey" not in html
    assert "Ley 16.744, Art. 68 inciso tres" in html
    assert "RECEPCIÓN DE Reglamento Interno de Orden, Higiene y Seguridad (RIOHS)" in html
    assert "desonocimiento de las normas contenidas en este reglamento" in html
    assert "RECEPCIÓN DE EVALUACIÓN INICIAL DE RIESGOS LABORALES (EIRL)" in html
    assert "Servicios Prueba SPA" in html
    assert "Firma / Timbre del Instructor o Supervisor" in html


def test_simulated_signatures_are_visible_in_source_signature_spaces():
    snapshot = _contract_snapshot()
    html = render_document_html(snapshot, {
        "worker_name": "Ana Pérez",
        "worker_signed_at": "2026-10-06T14:05:00+00:00",
        "supervisor_name": "Admin de Prueba",
        "supervisor_signed_at": "2026-10-06T14:10:00+00:00",
    })
    assert html.count("FIRMA SIMULADA · ENTORNO LOCAL") == 2
    assert "Ana Pérez" in html and "Admin de Prueba" in html
    assert "06 / 10 / 2026 14:05" in html and "06 / 10 / 2026 14:10" in html
    assert "Constancia de consentimiento electrónico" not in html
    assert "No corresponde a una firma electrónica" not in html


def test_epp_template_exposes_per_item_conformity_without_an_individual_signature():
    snapshot = {
        "title": "Registro EPP",
        "version": 1,
        "kind": "EPP",
        "prepared_at": "2026-10-06",
        "person": {"first_name": "Ana", "last_name": "Pérez", "rut": "11.111.111-1"},
        "employment": {},
        "event": None,
        "data": {},
        "epp": {
            "riohs_received": True,
            "eirl_completed": False,
            "items": [{"label": "Guantes", "received_at": "2026-10-01", "returned_at": None,
                       "received_conformity": True}],
        },
        "signature_notice": "Firma simulada local",
    }
    html = render_document_html(snapshot)
    assert "Guantes</td><td>01 / 10 / 2026</td><td class='conformity'>Sí" in html
    assert "RECEPCIÓN DE Reglamento Interno" in html and "EIRL" in html
    assert "Firma del Trabajador" in html


def test_epp_catalog_case_variants_fill_original_row_without_duplicate():
    snapshot = {
        "kind": "EPP",
        "person": {"first_name": "Ana", "last_name": "Pérez", "rut": "12.345.678-5"},
        "employment": {},
        "data": {},
        "epp": {
            "items": [{"label": "Bloqueador solar", "received_at": "2026-10-01", "received_conformity": True}]
        },
    }
    html = render_document_html(snapshot)
    assert html.count("Bloqueador solar</td>") == 1
    assert "Bloqueador solar</td><td>01 / 10 / 2026</td>" in html


@pytest.mark.parametrize(
    "snapshot,expected_text",
    [
        (_contract_snapshot(), "El presente contrato se firma en dos ejemplares"),
        (
            {
                "kind": "EPP",
                "person": {"first_name": "Ana", "last_name": "Pérez", "rut": "11.111.111-1"},
                "employment": {"position": "Técnica"},
                "employer_name": "Empresa de Prueba",
                "data": {},
                "epp": {"items": [{"label": "Guantes", "received_at": "2026-10-01", "received_conformity": True}]},
                "signature_notice": "Firma simulada local",
            },
            "RECEPCIÓN DE EVALUACIÓN INICIAL DE RIESGOS LABORALES",
        ),
    ],
)
def test_local_playwright_pdfs_are_valid_readable_and_contain_source_text(snapshot, expected_text):
    pdf, page_count = render_html(render_document_html(snapshot))
    reader = PdfReader(BytesIO(pdf))
    assert pdf.startswith(b"%PDF-")
    assert page_count == len(reader.pages) >= 1
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert expected_text in text
    assert len(hashlib.sha256(pdf).hexdigest()) == 64


def test_local_epp_playwright_pdf_contains_only_the_event_configured_items():
    snapshot = {
        "kind": "EPP",
        "person": {"first_name": "Ana", "last_name": "Pérez", "rut": "11.111.111-1"},
        "employment": {"position": "Técnica"},
        "employer_name": "Empresa de Prueba",
        "data": {},
        "epp": {"items": [{"label": "Guantes", "received_at": "2026-10-01", "returned_at": None, "received_conformity": True}]},
        "signature_notice": "Firma simulada local",
    }
    pdf, page_count = render_html(render_document_html(snapshot))
    reader = PdfReader(BytesIO(pdf))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert page_count == len(reader.pages) >= 1
    assert "Guantes" in text
    assert "Polera con filtro UV" not in text and "Protectores auditivos" not in text
    assert "Desconocimiento" not in text
    assert "Sí" in text
    assert "Guantes" in text and "Firma del Trabajador" in text


def test_expired_active_link_is_reported_without_mutating_document_status():
    now = datetime.now(UTC).replace(tzinfo=None)
    expired_request = SimpleNamespace(
        revoked_at=None,
        expires_at=now - timedelta(seconds=1),
        created_at=now - timedelta(days=8),
        worker_signed_at=None,
        opened_at=None,
    )
    doc = SimpleNamespace(
        status="SENT", signatures=[expired_request], id=uuid4(), person_id=uuid4(),
        person=SimpleNamespace(first_name="Ana", last_name="Pérez", rut="11.111.111-1"),
        employment_id=None, event_id=None, kind="CONTRACT", title="Contrato", frozen_snapshot=None,
        data={}, epp_delivery=None, version=1, created_at=now, updated_at=now,
        completed_at=None, original_pdf_key=None, signed_pdf_key=None,
    )
    assert _read(doc)["status"] == "EXPIRED"
    assert doc.status == "SENT"


def test_new_epp_version_copies_delivery_and_independent_items():
    source = SimpleNamespace(
        riohs_received=True,
        eirl_completed=True,
        delivery_date=date(2026, 10, 1),
        supervisor_id=uuid4(),
        items=[SimpleNamespace(
            item_code="GLOVES", label="Guantes", received_at=date(2026, 10, 1),
            returned_at=date(2026, 10, 3), received_conformity=True,
        )],
    )

    class FakeSession:
        added = None

        def add(self, value):
            self.added = value

        def flush(self):
            pass

    session = FakeSession()
    copied = _copy_epp_delivery(session, source, uuid4(), uuid4())
    assert copied is session.added
    assert copied.document_id is not None
    assert not copied.riohs_received and not copied.eirl_completed
    assert copied.delivery_date == source.delivery_date
    assert copied.supervisor_id == source.supervisor_id
    assert len(copied.items) == 1
    clone = copied.items[0]
    assert (clone.item_code, clone.label, clone.received_at, clone.returned_at) == (
        "GLOVES", "Guantes", date(2026, 10, 1), date(2026, 10, 3)
    )
    assert clone.received_conformity is False
    clone.label = "Versión nueva"
    assert source.items[0].label == "Guantes"


def test_worker_epp_acknowledgments_fill_receipt_date_and_conformity_in_preview():
    snapshot = {
        "kind": "EPP",
        "epp": {
            "delivery_date": None,
            "items": [
                {"item_code": "JOCKEY", "label": "Jockey", "received_at": None},
                {"item_code": "GLOVES", "label": "Guantes", "received_at": None},
            ],
        },
    }
    worker_ack = {
        "riohs_received": True,
        "eirl_completed": True,
        "received_item_codes": ["JOCKEY", "GLOVES"],
        "confirmed_at": "2026-10-07T12:30:00",
    }

    updated = _apply_worker_epp_acknowledgments(snapshot, worker_ack)
    html = render_document_html(updated)

    assert updated["epp"]["delivery_date"] == "2026-10-07"
    assert all(item["received_at"] == "2026-10-07" for item in updated["epp"]["items"])
    assert all(item["received_conformity"] for item in updated["epp"]["items"])
    assert "07 / 10 / 2026" in html
    assert "class='conformity'>" in html


def test_preview_pdf_returns_immutable_signed_pdf_when_available(monkeypatch):
    document_id = uuid4()
    document = SimpleNamespace(id=document_id, signed_pdf_key="documents/example/v2-signed.pdf")

    class Query:
        def where(self, *_args):
            return self

    class FakeSession:
        def scalar(self, _query):
            return document

    monkeypatch.setattr(
        "app.services.managed_documents_service._document_query", lambda: Query()
    )
    monkeypatch.setattr(
        "app.services.managed_documents_service._allowed", lambda *_args, **_kwargs: None
    )
    monkeypatch.setattr(
        "app.services.managed_documents_service._read_pdf",
        lambda key: b"%PDF-signed-artifact" if key == document.signed_pdf_key else b"",
    )

    assert preview_pdf(FakeSession(), document_id, SimpleNamespace()) == b"%PDF-signed-artifact"


def test_public_token_is_stored_as_one_way_digest():
    token = "sample-high-entropy-token-value"
    digest = token_digest(token)
    assert len(digest) == 64
    assert token not in digest


def test_signature_provider_issues_random_one_time_capability_with_expiration():
    now = datetime(2026, 10, 6)
    capability = LocalSimulatedSignatureProvider().create_signature_request(now, timedelta(days=7))
    assert len(capability.token) >= 40
    assert capability.token_hash == token_digest(capability.token)
    assert capability.token_hash != capability.token
    assert capability.expires_at == now + timedelta(days=7)


def test_local_document_storage_path_is_private_uuid_scoped_and_traversal_safe(monkeypatch):
    private_root = Path(__file__).resolve().parents[1] / "uploads" / "private"
    monkeypatch.setattr(settings, "local_private_storage_root", str(private_root))
    storage = LocalPrivateDocumentStorage()
    document_id = uuid4()
    path = storage._path(f"documents/{document_id}/v1-original.pdf")
    assert path == private_root.resolve() / "documents" / str(document_id) / "v1-original.pdf"
    with pytest.raises(HTTPException):
        storage._path("../../outside.pdf")


def test_document_storage_can_write_and_read_private_r2_objects(monkeypatch):
    stored = {}

    def fake_save(key, content, *, content_type):
        stored[key] = (content, content_type)
        return key

    def fake_read(key):
        return stored[key]

    monkeypatch.setattr(settings, "documents_use_r2_storage", True)
    monkeypatch.setattr(settings, "cloudflare_r2_bucket", "test-bucket")
    monkeypatch.setattr(settings, "cloudflare_r2_account_id", "test-account")
    monkeypatch.setattr(settings, "cloudflare_r2_access_key_id", "test-access")
    monkeypatch.setattr(settings, "cloudflare_r2_secret_access_key", "test-secret")
    monkeypatch.setattr(settings, "r2_private_prefix", "private")
    monkeypatch.setattr(
        "app.services.document_storage_service.save_private_object_to_r2", fake_save
    )
    monkeypatch.setattr(
        "app.services.document_storage_service.read_private_object_from_r2", fake_read
    )
    storage = ConfiguredPrivateDocumentStorage()
    document_id = uuid4()
    pdf = b"%PDF-1.7\n%%EOF"

    key, checksum = storage.store(document_id, 2, "signed", pdf)

    assert key == f"private/documents/{document_id}/v2-signed.pdf"
    assert checksum == hashlib.sha256(pdf).hexdigest()
    assert storage.read(key) == pdf
    assert stored[key] == (pdf, "application/pdf")


def test_document_storage_uses_new_immutable_key_when_previous_upload_was_orphaned(monkeypatch):
    from app.services.document_storage_service import CloudflareR2PrivateDocumentStorage

    stored = {}
    attempts = 0

    def fake_save(key, content, *, content_type):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            raise FileExistsError("orphaned object from a previous interrupted request")
        stored[key] = (content, content_type)
        return key

    monkeypatch.setattr(settings, "r2_private_prefix", "private")
    monkeypatch.setattr(
        "app.services.document_storage_service.save_private_object_to_r2", fake_save
    )
    storage = CloudflareR2PrivateDocumentStorage()
    document_id = uuid4()
    pdf = b"%PDF-1.7\n%%EOF"

    key, checksum = storage.store(document_id, 1, "original", pdf)

    assert attempts == 2
    assert key.startswith(f"private/documents/{document_id}/v1-original-")
    assert key.endswith(".pdf")
    assert checksum == hashlib.sha256(pdf).hexdigest()
    assert stored[key] == (pdf, "application/pdf")


def test_paper_print_batch_reprint_strips_saved_employer_signature(monkeypatch):
    from app.services.managed_documents_service import build_paper_print_batch

    event_id, person_id, sent_person_id = uuid4(), uuid4(), uuid4()
    pdf_writer = __import__("pypdf").PdfWriter()
    pdf_writer.add_blank_page(width=612, height=792)
    output = BytesIO()
    pdf_writer.write(output)
    pdf_bytes = output.getvalue()
    draft = SimpleNamespace(
        person_id=person_id, kind="CONTRACT", status="DRAFT", version=1,
        updated_at=datetime.now(UTC), original_pdf_key="draft.pdf",
        signature_method="PAPER", paper_status="PREPARED", paper_batch_items=[object()],
        event_id=event_id, client_id=uuid4(), id=uuid4(),
        frozen_snapshot={"kind": "CONTRACT", "employer_signature": {
            "signer_name": "Representante", "signature_data_url": "data:image/png;base64,secret",
        }},
        person=SimpleNamespace(first_name="Ana", last_name="Pérez"),
    )
    sent = SimpleNamespace(
        person_id=sent_person_id, kind="CONTRACT", status="SENT", version=1,
        updated_at=datetime.now(UTC), original_pdf_key="sent.pdf",
        signature_method="ELECTRONIC", paper_status=None, paper_batch_items=[],
        person=SimpleNamespace(first_name="Luis", last_name="Rojas"),
    )

    event = SimpleNamespace(name="Faena de prueba", client_id=draft.client_id, id=event_id)

    class FakeDb:
        def get(self, model, key):
            return event

        def scalars(self, statement):
            return iter([draft, sent])

        def add(self, _item):
            pass

        def commit(self):
            pass

    monkeypatch.setattr(managed_documents_service, "can_manage_event", lambda *_: True)
    monkeypatch.setattr(managed_documents_service, "_allowed", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(managed_documents_service, "_raise_if_document_incomplete", lambda *_args, **_kwargs: None)
    def unexpected_signature_profile(*_args, **_kwargs):
        raise AssertionError("paper flow must not need employer signature")

    monkeypatch.setattr(managed_documents_service, "_signature_for_event", unexpected_signature_profile)
    monkeypatch.setattr(
        managed_documents_service,
        "_read_pdf",
        lambda *_args: (_ for _ in ()).throw(AssertionError("must regenerate instead of reuse a stale PDF")),
    )
    rendered = {}
    monkeypatch.setattr(
        managed_documents_service,
        "render_document_html",
        lambda snapshot, **kwargs: rendered.update(snapshot=snapshot, kwargs=kwargs) or "paper html",
    )
    monkeypatch.setattr(managed_documents_service, "render_html", lambda _html: (pdf_bytes, 1))
    monkeypatch.setattr(managed_documents_service.document_storage, "store", lambda *_args: ("paper-batch.pdf", "checksum"))

    pdf, batch = build_paper_print_batch(
        FakeDb(), event_id, [person_id, sent_person_id], ["CONTRACT"], SimpleNamespace(id=uuid4()),
        allow_reprint_existing=True,
    )

    assert batch.document_count == 1
    assert batch.page_count == 1
    assert len(PdfReader(BytesIO(pdf)).pages) == 1
    assert draft.status == "DRAFT"
    assert sent.status == "SENT"
    assert draft.signature_method == "PAPER"
    assert draft.paper_status == "PREPARED"
    assert "employer_signature" not in draft.frozen_snapshot
    assert "include_employer_signature" not in rendered["kwargs"]
    assert sent.signature_method == "ELECTRONIC"


def test_paper_print_batch_request_rejects_duplicate_or_unknown_kinds():
    from app.schemas.document_schema import PaperBatchPrintRequest

    with pytest.raises(ValueError):
        PaperBatchPrintRequest(person_ids=[uuid4()], kinds=["CONTRACT", "CONTRACT"])
    with pytest.raises(ValueError):
        PaperBatchPrintRequest(person_ids=[uuid4()], kinds=["UNKNOWN"])

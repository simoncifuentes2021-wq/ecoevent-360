"""Service-level local workflows using real PostgreSQL and rendered PDFs."""

from datetime import UTC, datetime, timedelta
import base64
from io import BytesIO
from uuid import uuid4

import pytest
from PIL import Image, ImageDraw
from pypdf import PdfReader
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.session import set_rls_context
from app.models.audit_log import AuditLog
from app.models.core import Client, Event, User
from app.models.documents import DocumentEmployment, DocumentPerson, DocumentSignatureRequest, EppDelivery, EppDeliveryItem, ManagedDocument
from app.models.enums import EventStatus, UserRole
from app.schemas.document_schema import (
    DocumentCreate,
    DocumentUpdate,
    EppInput,
    EppItemInput,
    EppWorkerAcknowledgement,
    DocumentContractDefaults,
    ExistingPeopleAssignmentCommit,
    ExistingPersonAssignmentInput,
    GenerateMissingDraftsCommit,
    EmploymentInput,
    EmploymentUpdate,
    PersonImportCommit,
    PersonImportRow,
    PersonCreate,
    PersonUpdate,
    RepresentativeSignatureSave,
)
from app.services import document_batch_service, document_people_service, document_reporting_service, managed_documents_service
from app.services.audit_log_service import create_audit_log
from app.api.routers.documents import get_person_document_history


def _signature_data_url() -> str:
    image = Image.new("RGBA", (360, 120), (255, 255, 255, 0))
    ImageDraw.Draw(image).line([(25, 90), (90, 30), (140, 85), (205, 25), (270, 70), (330, 35)], fill=(15, 35, 30, 255), width=5)
    output = BytesIO()
    image.save(output, format="PNG")
    return "data:image/png;base64," + base64.b64encode(output.getvalue()).decode("ascii")


def _configure_representative_signature(db, staff):
    managed_documents_service.save_contract_defaults(
        db,
        DocumentContractDefaults(
            employer_name="Greenway Servicios Ambientales SpA",
            employer_rut="76.747.233-1",
            representative_name="Representante",
            representative_rut="17.188.611-2",
            employer_address="Santiago",
            employer_commune="Santiago",
        ),
        staff,
    )
    managed_documents_service.save_representative_signature(
        db, RepresentativeSignatureSave(signature_data_url=_signature_data_url()), staff
    )


@pytest.fixture
def local_document_staff(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "local_private_storage_root", str(tmp_path / "private"))
    monkeypatch.setattr(settings, "documents_use_r2_storage", False)
    engine = create_engine(settings.database_url, connect_args={"prepare_threshold": None})
    db = Session(engine, expire_on_commit=False)
    suffix = uuid4().hex
    client = Client(business_name=f"Workflow {suffix}")
    db.add(client)
    db.flush()
    user = User(
        client_id=client.id,
        full_name="Admin Workflow",
        email=f"workflow-{suffix}@example.test",
        password_hash="test-only",
        role=UserRole.ADMIN,
    )
    db.add(user)
    db.flush()
    event = Event(
        client_id=client.id,
        created_by=user.id,
        name=f"Evento workflow {suffix}",
        start_date=datetime.now(UTC).replace(tzinfo=None),
        end_date=(datetime.now(UTC) + timedelta(days=1)).replace(tzinfo=None),
        status=EventStatus.IN_PROGRESS,
    )
    db.add(event)
    db.commit()
    set_rls_context(db, user_id=user.id, role=user.role, client_id=client.id)
    try:
        yield db, client, user, event
    finally:
        db.close()
        engine.dispose()


def _person_for_staff(db, client, staff, *, event_id=None):
    digits = str(client.id.int % 100_000_000).zfill(8)
    total, multiplier = 0, 2
    for digit in reversed(digits):
        total += int(digit) * multiplier
        multiplier = 2 if multiplier == 7 else multiplier + 1
    check = 11 - total % 11
    verifier = "0" if check == 11 else "K" if check == 10 else str(check)
    return document_people_service.create_person(
        db,
        PersonCreate(
            client_id=client.id,
            first_name="Ana",
            last_name="Pérez",
            rut=f"{digits}-{verifier}",
            email=f"worker-{digits}@example.com",
            address="Calle 1",
            commune="Santiago",
            employment=EmploymentInput(event_id=event_id, position="Técnica", workplace="Recinto"),
        ),
        staff,
    )


def _valid_rut_from_seed(seed: int) -> str:
    digits = str(seed % 100_000_000).zfill(8)
    total, multiplier = 0, 2
    for digit in reversed(digits):
        total += int(digit) * multiplier
        multiplier = 2 if multiplier == 7 else multiplier + 1
    check = 11 - total % 11
    verifier = "0" if check == 11 else "K" if check == 10 else str(check)
    return f"{digits}-{verifier}"


def _pdf_text(content):
    reader = PdfReader(BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def test_person_profile_and_faena_assignment_are_edited_separately(local_document_staff):
    db, client, staff, event = local_document_staff
    person = _person_for_staff(db, client, staff, event_id=event.id)
    assignment = db.scalar(
        select(DocumentEmployment).where(
            DocumentEmployment.person_id == person.id,
            DocumentEmployment.event_id == event.id,
        )
    )

    updated_person = document_people_service.update_person(
        db,
        person.id,
        PersonUpdate(
            first_name="Ana María",
            last_name="Pérez",
            rut=person.rut,
            email="ana@example.com",
            phone="+56912345678",
            address="Calle 2",
            commune="Temuco",
        ),
        staff,
        event_id=event.id,
    )
    updated_assignment = document_people_service.update_employment(
        db,
        person.id,
        assignment.id,
        EmploymentUpdate(
            position="Supervisora",
            area="Operaciones",
            workplace="Recinto B",
            start_date=None,
            end_date=None,
            remuneration="75000",
        ),
        staff,
        event_id=event.id,
    )

    assert updated_person.first_name == "Ana María"
    assert updated_person.address == "Calle 2"
    assert updated_assignment.event_id == event.id
    assert updated_assignment.position == "Supervisora"
    assert str(updated_assignment.remuneration) == "75000.00"


def test_general_contract_preview_prepare_link_worker_consent_completion_and_email(local_document_staff, monkeypatch):
    db, client, staff, event = local_document_staff
    _configure_representative_signature(db, staff)
    person = _person_for_staff(db, client, staff, event_id=event.id)
    assignment = db.scalar(select(DocumentEmployment).where(DocumentEmployment.person_id == person.id))
    assignment.start_date = datetime.now(UTC).date()
    assignment.end_date = (datetime.now(UTC) + timedelta(days=1)).date()
    assignment.remuneration = 75000
    person.first_name = "sIMON"
    person.last_name = "cIFUENTES"
    db.commit()
    created = managed_documents_service.create_document(
        db,
        DocumentCreate(
            person_id=person.id,
            event_id=event.id,
            kind="CONTRACT",
            title="Contrato general local",
            data={
                "contract_date": datetime.now(UTC).date().isoformat(),
                "employer_name": "Greenway Servicios Ambientales SpA",
                "employer_rut": "76.747.233-1",
                "representative_name": "Representante",
                "representative_rut": "17.188.611-2",
                "employer_address": "Santiago",
                "workplace": "Recinto",
            },
        ),
        staff,
    )
    document_id = created["id"]
    assert created["event_id"] == event.id
    preview = managed_documents_service.preview_pdf(db, document_id, staff)
    assert _pdf_text(preview)
    prepared = managed_documents_service.prepare_document(db, document_id, staff)
    assert prepared["status"] == "READY"
    link_emails = []
    monkeypatch.setattr(
        managed_documents_service,
        "send_document_signature_link_email",
        lambda **kwargs: (link_emails.append(kwargs) or True),
    )
    sent = managed_documents_service.send_for_signature(db, document_id, staff)
    assert sent["email_sent"] is True
    assert sent["recipient"] == person.email
    assert len(link_emails) == 1
    assert link_emails[0]["portal_url"] == sent["link"]["url"]
    assert link_emails[0]["recipient"] == person.email
    token = sent["link"]["url"].split("#", 1)[1]
    opened = managed_documents_service.open_public_document(db, token, person.rut)
    assert opened["status"] == "VIEWED"
    assert opened["person_name"] == "Simon Cifuentes"
    assert opened["snapshot"]["employer_signature"]["signer_name"] == "Representante"
    assert "signature_data_url" not in opened["snapshot"]["employer_signature"]
    assert db.get(ManagedDocument, document_id).frozen_snapshot["employer_signature"]["signature_data_url"].startswith("data:image/png;base64,")
    public_pdf = managed_documents_service.preview_public_document_pdf(db, token, person.rut)
    assert "CONTRATO POR OBRA O FAENA" in _pdf_text(public_pdf)
    with pytest.raises(Exception) as consent_error:
        managed_documents_service.sign_public_document(
            db, token, person.rut, "Simon Cifuentes", True, False,
            worker_signature_data_url=_signature_data_url(),
        )
    assert getattr(consent_error.value, "status_code", None) == 422
    email_calls = []
    monkeypatch.setattr(
        managed_documents_service,
        "send_signed_document_copy_email",
        lambda **kwargs: (email_calls.append(kwargs) or True),
    )
    worker_result = managed_documents_service.sign_public_document(
        db, token, person.rut, "Simon Cifuentes", True, True,
        worker_signature_data_url=_signature_data_url(),
        client_ip="203.0.113.10",
        user_agent="Test browser",
    )
    assert worker_result["status"] == "COMPLETED"
    assert worker_result["email_sent"] is True
    assert len(email_calls) == 1 and email_calls[0]["recipient"] == person.email
    assert email_calls[0]["pdf_content"].startswith(b"%PDF-")
    assert person.first_name == "sIMON" and person.last_name == "cIFUENTES"
    completed = managed_documents_service.get_document(db, document_id, staff)
    assert completed["status"] == "COMPLETED"
    original = managed_documents_service.download_document(db, document_id, staff)
    signed = managed_documents_service.download_document(db, document_id, staff, signed=True)
    assert original.startswith(b"%PDF-") and signed.startswith(b"%PDF-")
    assert original != signed
    assert "Firma manuscrita del trabajador" in _pdf_text(signed)
    assert "Firma del empleador / representante legal" in _pdf_text(signed)
    assert "Constancia de consentimiento electrónico" not in _pdf_text(signed)
    assert "No corresponden a firma electrónica avanzada" not in _pdf_text(signed)
    request = db.scalar(select(DocumentSignatureRequest).where(DocumentSignatureRequest.document_id == document_id))
    assert request.worker_form_data["electronic_consent"] is True
    assert request.worker_form_data["electronic_consent_version"]
    assert request.worker_form_data["electronic_consent_at"].endswith("Z")
    assert request.worker_form_data["electronic_consent_text"]
    assert request.worker_form_data["client_ip"] == "203.0.113.10"
    assert request.worker_form_data["signed_pdf_sha256"] == db.get(ManagedDocument, document_id).signed_sha256
    actions = set(db.scalars(select(AuditLog.action).where(
        AuditLog.entity_id == document_id,
        AuditLog.action.in_(("DOCUMENT_PUBLIC_OPENED", "DOCUMENT_WORKER_SIGNED")),
    )))
    assert actions == {"DOCUMENT_PUBLIC_OPENED", "DOCUMENT_WORKER_SIGNED"}


def test_documents_sent_together_email_both_pdfs_only_after_both_signatures(
    local_document_staff, monkeypatch
):
    db, client, staff, event = local_document_staff
    _configure_representative_signature(db, staff)
    person = _person_for_staff(db, client, staff, event_id=event.id)
    assignment = db.scalar(
        select(DocumentEmployment).where(DocumentEmployment.person_id == person.id)
    )
    assignment.start_date = datetime.now(UTC).date()
    assignment.end_date = (datetime.now(UTC) + timedelta(days=1)).date()
    assignment.remuneration = 75000
    db.commit()
    contract = managed_documents_service.create_document(
        db,
        DocumentCreate(
            person_id=person.id,
            event_id=event.id,
            kind="CONTRACT",
            title="Contrato de faena",
            data={
                "contract_date": datetime.now(UTC).date().isoformat(),
                "employer_name": "Greenway Servicios Ambientales SpA",
                "employer_rut": "76.747.233-1",
                "representative_name": "Representante",
                "representative_rut": "17.188.611-2",
                "employer_address": "Santiago",
                "workplace": "Recinto",
            },
        ),
        staff,
    )
    epp = managed_documents_service.create_document(
        db,
        DocumentCreate(
            person_id=person.id,
            event_id=event.id,
            kind="EPP",
            title="Registro EPP",
            epp=EppInput(items=[EppItemInput(
                item_code="GLOVES",
                label="Guantes",
                received_at=datetime.now(UTC).date(),
                received_conformity=True,
            )]),
        ),
        staff,
    )
    monkeypatch.setattr(managed_documents_service, "send_document_pair_email", lambda **_: True)
    single_link_emails = []
    monkeypatch.setattr(
        managed_documents_service,
        "send_document_signature_link_email",
        lambda **kwargs: (single_link_emails.append(kwargs) or True),
    )
    bundle_emails = []
    monkeypatch.setattr(
        managed_documents_service,
        "send_signed_document_bundle_email",
        lambda **kwargs: (bundle_emails.append(kwargs) or True),
    )

    sent = managed_documents_service.send_documents_together(
        db, contract["id"], epp["id"], staff
    )
    assert sent["email_sent"] is True
    assert single_link_emails == []
    assert len(sent["links"]) == 1
    bundle_parts = sent["links"][0]["url"].split("#", 1)[1].split(".")
    assert bundle_parts[0] == "bundle"
    contract_token, epp_token = bundle_parts[1:]
    managed_documents_service.open_public_document(db, contract_token, person.rut)
    managed_documents_service.open_public_document(db, epp_token, person.rut)

    first_result = managed_documents_service.sign_public_document(
        db, contract_token, person.rut, "Ana Pérez", True, True,
        worker_signature_data_url=_signature_data_url(),
    )
    assert first_result["email_sent"] is False
    assert "Cuando completes la firma del otro documento" in first_result["delivery_error"]
    assert bundle_emails == []

    second_result = managed_documents_service.sign_public_document(
        db, epp_token, person.rut, "Ana Pérez", True, True,
        EppWorkerAcknowledgement(
            riohs_received=True,
            eirl_completed=True,
            items_reviewed=True,
            received_item_codes=["GLOVES"],
        ),
        worker_signature_data_url=_signature_data_url(),
    )
    assert second_result["email_sent"] is True
    assert len(bundle_emails) == 1
    assert bundle_emails[0]["recipient"] == person.email
    assert len(bundle_emails[0]["attachments"]) == 2
    assert all(item["content"].startswith(b"%PDF-") for item in bundle_emails[0]["attachments"])
    requests = db.scalars(
        select(DocumentSignatureRequest).where(
            DocumentSignatureRequest.document_id.in_([contract["id"], epp["id"]]),
            DocumentSignatureRequest.bundle_id.is_not(None),
        )
    ).all()
    assert len(requests) == 2
    assert all(item.worker_form_data["signed_copy_email_status"] == "SENT" for item in requests)


def test_incomplete_draft_can_be_saved_but_cannot_be_prepared(local_document_staff):
    db, client, staff, event = local_document_staff
    person = _person_for_staff(db, client, staff, event_id=event.id)
    created = managed_documents_service.create_document(
        db,
        DocumentCreate(person_id=person.id, event_id=event.id, kind="CONTRACT", title="Borrador incompleto", data={}),
        staff,
    )
    assert created["status"] == "DRAFT"
    with pytest.raises(Exception) as error:
        managed_documents_service.prepare_document(db, created["id"], staff)
    assert getattr(error.value, "status_code", None) == 422
    assert "RUT del empleador" in str(error.value.detail)
    assert managed_documents_service.get_document(db, created["id"], staff)["status"] == "DRAFT"


def test_incomplete_ready_document_cannot_be_sent_and_empty_epp_cannot_be_prepared(local_document_staff):
    db, client, staff, event = local_document_staff
    person = _person_for_staff(db, client, staff, event_id=event.id)
    contract = managed_documents_service.create_document(
        db,
        DocumentCreate(person_id=person.id, event_id=event.id, kind="CONTRACT", title="Incompleto", data={}),
        staff,
    )
    contract_model = db.get(ManagedDocument, contract["id"])
    contract_model.status = "READY"
    db.commit()
    with pytest.raises(Exception) as send_error:
        managed_documents_service.send_for_signature(db, contract["id"], staff)
    assert getattr(send_error.value, "status_code", None) == 422
    assert db.scalar(select(DocumentSignatureRequest.id).where(DocumentSignatureRequest.document_id == contract["id"])) is None
    assert managed_documents_service.get_document(db, contract["id"], staff)["status"] == "READY"

    epp = managed_documents_service.create_document(
        db,
        DocumentCreate(person_id=person.id, event_id=event.id, kind="EPP", title="EPP sin elementos", epp=EppInput(items=[])),
        staff,
    )
    delivery = db.scalar(select(EppDelivery).where(EppDelivery.document_id == epp["id"]))
    for item in db.scalars(select(EppDeliveryItem).where(EppDeliveryItem.delivery_id == delivery.id)).all():
        db.delete(item)
    db.commit()
    with pytest.raises(Exception) as epp_error:
        managed_documents_service.prepare_document(db, epp["id"], staff)
    assert getattr(epp_error.value, "status_code", None) == 422
    assert "Al menos un elemento EPP" in str(epp_error.value.detail)
    assert managed_documents_service.get_document(db, epp["id"], staff)["status"] == "DRAFT"


def test_event_epp_worker_supervisor_and_new_version_are_independent(local_document_staff):
    db, client, staff, event = local_document_staff
    _configure_representative_signature(db, staff)
    person = _person_for_staff(db, client, staff, event_id=event.id)
    created = managed_documents_service.create_document(
        db,
        DocumentCreate(
            person_id=person.id,
            event_id=event.id,
            kind="EPP",
            title="Registro EPP local",
            epp=EppInput(
                riohs_received=True,
                eirl_completed=True,
                items=[EppItemInput(
                    item_code="GLOVES",
                    label="Guantes de trabajo",
                    received_at=datetime.now(UTC).date(),
                    received_conformity=True,
                )],
            ),
        ),
        staff,
    )
    document_id = created["id"]
    assert created["event_id"] == event.id
    preview = managed_documents_service.preview_pdf(db, document_id, staff)
    assert "RECIBÍ CONFORME" in _pdf_text(preview)
    prepared = managed_documents_service.prepare_document(db, document_id, staff)
    assert prepared["status"] == "READY"
    sent = managed_documents_service.send_for_signature(db, document_id, staff)
    token = sent["link"]["url"].split("#", 1)[1]
    managed_documents_service.open_public_document(db, token, person.rut)
    worker_pdf = managed_documents_service.preview_public_document_pdf(db, token, person.rut)
    assert "Guantes de trabajo" in _pdf_text(worker_pdf)
    managed_documents_service.sign_public_document(
        db,
        token,
        person.rut,
        "Ana Pérez",
        True,
        True,
        EppWorkerAcknowledgement(
            riohs_received=True,
            eirl_completed=True,
                items_reviewed=True,
            received_item_codes=["GLOVES"],
        ),
        worker_signature_data_url=_signature_data_url(),
    )
    completed = managed_documents_service.get_document(db, document_id, staff)
    assert completed["status"] == "COMPLETED"
    signed_pdf = managed_documents_service.download_document(db, document_id, staff, signed=True)
    assert "Firma del Trabajador" in _pdf_text(signed_pdf)
    assert "Firma / Timbre del Instructor o Supervisor" in _pdf_text(signed_pdf)

    next_version = managed_documents_service.create_new_version(db, document_id, staff)
    assert next_version["version"] == 2 and next_version["status"] == "DRAFT"
    assert next_version["epp"]["riohs_received"] is False
    assert next_version["epp"]["eirl_completed"] is False
    cloned_item = next(item for item in next_version["epp"]["items"] if item["item_code"] == "GLOVES")
    assert cloned_item["received_conformity"] is False
    managed_documents_service.update_document(
        db,
        next_version["id"],
        DocumentUpdate(epp=EppInput(
            riohs_received=False,
            eirl_completed=True,
            items=[EppItemInput(item_code="GLOVES", label="Guantes versión 2")],
        )),
        staff,
    )
    old = managed_documents_service.get_document(db, document_id, staff)
    updated = managed_documents_service.get_document(db, next_version["id"], staff)
    old_item = next(item for item in old["epp"]["items"] if item["item_code"] == "GLOVES")
    new_item = next(item for item in updated["epp"]["items"] if item["item_code"] == "GLOVES")
    assert old_item["label"] == "Guantes de trabajo" and old_item["received_conformity"] is True
    assert new_item["label"] == "Guantes versión 2" and new_item["received_conformity"] is False


def test_batch_import_creates_event_document_pairs_once(local_document_staff, monkeypatch):
    db, client, staff, event = local_document_staff
    person = _person_for_staff(db, client, staff, event_id=event.id)
    monkeypatch.setattr(
        document_batch_service,
        "get_contract_defaults",
        lambda *_args: {
            "employer_name": "Servicios Greenway SPA",
            "employer_rut": "76.747.233-1",
            "representative_name": "Gonzalo Ormeño Bahamondes",
            "representative_rut": "17.188.611-2",
            "employer_address": "Ñuble N°1034",
            "employer_commune": "Santiago",
        },
    )
    payload = PersonImportCommit(
        client_id=client.id,
        event_id=event.id,
        rows=[
            PersonImportRow(
                first_name=person.first_name,
                last_name=person.last_name,
                rut=person.rut,
                position="Técnica",
            )
        ],
    )

    normalized_rut = document_people_service.normalize_rut(person.rut)
    assert person.rut_normalized == normalized_rut
    assert db.scalar(select(DocumentPerson.id).where(DocumentPerson.rut_normalized == normalized_rut)) == person.id
    preview = document_batch_service.preview_import_draft_pairs(db, payload, staff)
    assert preview["documents_to_create"] == 2
    assert preview["documents_existing"] == 0
    assert preview["results"][0]["person_action"] == "Actualizar ficha"
    assert db.scalar(select(ManagedDocument.id).where(ManagedDocument.person_id == person.id)) is None

    first = document_batch_service.generate_draft_pairs(db, payload, staff)
    assert first["workers_processed"] == 1
    assert first["contracts_created"] == first["epp_documents_created"] == 1
    assert first["documents_created"] == 2
    assert first["setup_warnings"] == []

    docs = db.scalars(
        select(ManagedDocument).where(
            ManagedDocument.person_id == person.id,
            ManagedDocument.event_id == event.id,
        )
    ).all()
    assert {doc.kind for doc in docs} == {"CONTRACT", "EPP"}
    assert all(doc.status == "DRAFT" and doc.employment_id is not None for doc in docs)
    contract = next(doc for doc in docs if doc.kind == "CONTRACT")
    assert contract.data["employer_name"] == "Servicios Greenway SPA"
    assert contract.data["duties"] == "Técnica"

    after_preview = document_batch_service.preview_import_draft_pairs(db, payload, staff)
    assert after_preview["documents_to_create"] == 0
    assert after_preview["documents_existing"] == 2

    repeated = document_batch_service.generate_draft_pairs(db, payload, staff)
    assert repeated["documents_created"] == 0
    assert repeated["documents_skipped"] == 2


def test_batch_import_uses_excel_duties_in_contract_and_saves_them_on_assignment(local_document_staff, monkeypatch):
    db, client, staff, event = local_document_staff
    person = _person_for_staff(db, client, staff, event_id=event.id)
    monkeypatch.setattr(document_batch_service, "get_contract_defaults", lambda *_args: {
        "employer_name": "Servicios Greenway SPA", "employer_rut": "76.747.233-1",
        "representative_name": "Gonzalo Ormeño Bahamondes", "representative_rut": "17.188.611-2",
        "employer_address": "Ñuble N°1034", "employer_commune": "Santiago",
    })
    duties = "Staff de aseo, limpieza y reciclaje"
    payload = PersonImportCommit(
        client_id=client.id,
        event_id=event.id,
        rows=[PersonImportRow(
            first_name=person.first_name, last_name=person.last_name, rut=person.rut,
            position="Staff", duties=duties,
        )],
    )

    result = document_batch_service.generate_draft_pairs(db, payload, staff)
    assert result["documents_created"] == 2
    contract = db.scalar(select(ManagedDocument).where(
        ManagedDocument.person_id == person.id,
        ManagedDocument.event_id == event.id,
        ManagedDocument.kind == "CONTRACT",
    ))
    assignment = db.get(DocumentEmployment, contract.employment_id)
    assert contract.data["duties"] == duties
    assert assignment.duties == duties


def test_batch_import_reports_row_failure_and_can_retry_only_that_worker(local_document_staff, monkeypatch):
    db, client, staff, event = local_document_staff
    first_rut = _valid_rut_from_seed(uuid4().int)
    failed_rut = _valid_rut_from_seed(uuid4().int)
    rows = [
        PersonImportRow(first_name="Ana", last_name="Uno", rut=first_rut, position="Operaria"),
        PersonImportRow(first_name="Luis", last_name="Dos", rut=failed_rut, position="Operario"),
    ]
    payload = PersonImportCommit(client_id=client.id, event_id=event.id, rows=rows)
    monkeypatch.setattr(document_batch_service, "get_contract_defaults", lambda *_args: {
        "employer_name": "Greenway", "employer_rut": "76.747.233-1",
        "representative_name": "Representante", "representative_rut": "17.188.611-2",
        "employer_address": "Santiago", "employer_commune": "Santiago",
    })
    original_generate = document_batch_service._generate_one_import_row
    failed_normalized = failed_rut.replace("-", "")

    def fail_second_row_once(db, row, normalized, *args):
        if normalized == failed_normalized:
            raise RuntimeError("error transitorio de prueba")
        return original_generate(db, row, normalized, *args)

    monkeypatch.setattr(document_batch_service, "_generate_one_import_row", fail_second_row_once)
    first_attempt = document_batch_service.generate_draft_pairs(db, payload, staff)
    assert first_attempt["workers_processed"] == 2
    assert first_attempt["workers_failed"] == 1
    assert [row["status"] for row in first_attempt["results"]] == ["SUCCESS", "ERROR"]
    assert first_attempt["documents_created"] == 2

    monkeypatch.setattr(document_batch_service, "_generate_one_import_row", original_generate)
    retry_payload = PersonImportCommit(client_id=client.id, event_id=event.id, rows=[rows[1]])
    retry = document_batch_service.generate_draft_pairs(db, retry_payload, staff)
    assert retry["workers_failed"] == 0
    assert retry["results"][0]["status"] == "SUCCESS"
    assert retry["documents_created"] == 2
    assert db.scalar(select(DocumentPerson.id).where(DocumentPerson.rut_normalized == failed_normalized)) is not None


def test_people_only_import_reports_each_row_and_retries_failed_row(local_document_staff, monkeypatch):
    db, client, staff, event = local_document_staff
    good = PersonImportRow(first_name="Nora", last_name="Lista", rut=_valid_rut_from_seed(uuid4().int), position="Operaria")
    failed = PersonImportRow(first_name="Pablo", last_name="Reintento", rut=_valid_rut_from_seed(uuid4().int), position="Operario")
    payload = PersonImportCommit(client_id=client.id, event_id=event.id, rows=[good, failed])
    original_employment = document_people_service._employment

    def fail_second_assignment(person, client_id, row, event_id=None):
        if row.first_name == "Pablo":
            raise RuntimeError("error transitorio de prueba")
        return original_employment(person, client_id, row, event_id)

    monkeypatch.setattr(document_people_service, "_employment", fail_second_assignment)
    first = document_people_service.commit_import(db, payload, staff)
    assert first["workers_failed"] == 1
    assert [row["status"] for row in first["results"]] == ["SUCCESS", "ERROR"]
    assert first["created"] == first["assignments"] == 1

    monkeypatch.setattr(document_people_service, "_employment", original_employment)
    retry = document_people_service.commit_import(
        db,
        PersonImportCommit(client_id=client.id, event_id=event.id, rows=[failed]),
        staff,
    )
    assert retry["workers_failed"] == 0
    assert retry["results"][0]["status"] == "SUCCESS"
    assert retry["created"] == retry["assignments"] == 1


def test_generate_missing_drafts_only_creates_absent_kinds(local_document_staff, monkeypatch):
    db, client, staff, event = local_document_staff
    person = _person_for_staff(db, client, staff, event_id=event.id)
    monkeypatch.setattr(document_batch_service, "get_contract_defaults", lambda *_args: {"employer_name": "Greenway"})
    monkeypatch.setattr(
        document_batch_service,
        "get_event_epp_template",
        lambda *_args: {"items": [{"item_code": "GLOVES", "label": "Guantes"}]},
    )
    assignment = db.scalar(
        select(DocumentEmployment).where(
            DocumentEmployment.person_id == person.id,
            DocumentEmployment.event_id == event.id,
        )
    )
    db.add(ManagedDocument(
        client_id=client.id,
        person_id=person.id,
        employment_id=assignment.id,
        event_id=event.id,
        created_by=staff.id,
        kind="CONTRACT",
        title="Contrato existente",
        data={"duties": "No modificar"},
    ))
    db.commit()

    payload = GenerateMissingDraftsCommit(event_id=event.id, person_ids=[person.id])
    result = document_batch_service.generate_missing_draft_pairs(db, payload, staff)
    assert result["contracts_created"] == 0
    assert result["epp_documents_created"] == 1
    assert result["results"][0]["created"] == ["EPP"]
    assert result["results"][0]["already_had"] == ["CONTRACT"]

    repeated = document_batch_service.generate_missing_draft_pairs(db, payload, staff)
    assert repeated["documents_created"] == 0
    assert repeated["documents_skipped"] == 2


def test_missing_drafts_preview_rejects_people_not_assigned_to_event(local_document_staff):
    db, client, staff, event = local_document_staff
    normalized_rut = document_people_service.normalize_rut(_valid_rut_from_seed(uuid4().int))
    unattached = DocumentPerson(
        first_name="Persona",
        last_name="Sin Faena",
        rut=normalized_rut,
        rut_normalized=normalized_rut,
    )
    db.add(unattached)
    db.commit()
    payload = GenerateMissingDraftsCommit(event_id=event.id, person_ids=[unattached.id])
    with pytest.raises(Exception) as preview_error:
        document_batch_service.preview_missing_draft_pairs(db, payload, staff)
    assert getattr(preview_error.value, "status_code", None) == 422
    with pytest.raises(Exception) as generate_error:
        document_batch_service.generate_missing_draft_pairs(db, payload, staff)
    assert getattr(generate_error.value, "status_code", None) == 422


def test_same_greenway_worker_is_reused_for_another_client_faena(local_document_staff):
    db, client_a, staff, event_a = local_document_staff
    person_a = _person_for_staff(db, client_a, staff, event_id=event_a.id)
    client_b = Client(business_name="Cliente contratante Superman")
    db.add(client_b)
    db.flush()
    event_b = Event(
        client_id=client_b.id,
        created_by=staff.id,
        name="Faena Superman",
        start_date=datetime.now(UTC).replace(tzinfo=None),
        end_date=(datetime.now(UTC) + timedelta(days=1)).replace(tzinfo=None),
        status=EventStatus.IN_PROGRESS,
    )
    db.add(event_b)
    db.commit()
    person_b = document_people_service.create_person(
        db,
        PersonCreate(
            client_id=client_b.id,
            first_name="Ana",
            last_name="Pérez",
            rut=person_a.rut,
            employment=EmploymentInput(event_id=event_b.id, position="Coordinadora"),
        ),
        staff,
    )
    assert person_b.id == person_a.id
    assignments = db.scalars(
        select(DocumentEmployment).where(DocumentEmployment.person_id == person_a.id)
    ).all()
    assert {assignment.event_id for assignment in assignments} == {event_a.id, event_b.id}

    second_doc = managed_documents_service.create_document(
        db,
        DocumentCreate(
            person_id=person_a.id,
            event_id=event_b.id,
            kind="CONTRACT",
            title="Contrato faena Superman",
            data={"employer_name": "Greenway App"},
        ),
        staff,
    )
    assert second_doc["event_id"] == event_b.id
    snapshot = managed_documents_service._snapshot(
        db, db.get(ManagedDocument, second_doc["id"])
    )
    assert snapshot["employer_name"] == "Greenway App"
    assert snapshot["event"]["name"] == "Faena Superman"


def test_assign_existing_person_to_another_faena_skips_already_assigned_and_preserves_history(local_document_staff):
    db, client, staff, source_event = local_document_staff
    person = _person_for_staff(db, client, staff, event_id=source_event.id)
    old_doc = managed_documents_service.create_document(
        db,
        DocumentCreate(
            person_id=person.id,
            event_id=source_event.id,
            kind="CONTRACT",
            title="Contrato anterior",
            data={"duties": "Cargo anterior"},
        ),
        staff,
    )
    target_event = Event(
        client_id=client.id,
        created_by=staff.id,
        name="Faena de destino",
        start_date=datetime.now(UTC).replace(tzinfo=None),
        end_date=(datetime.now(UTC) + timedelta(days=5)).replace(tzinfo=None),
        status=EventStatus.PLANNING,
    )
    db.add(target_event)
    db.commit()

    payload = ExistingPeopleAssignmentCommit(
        source_event_id=source_event.id,
        target_event_id=target_event.id,
        generate_documents=True,
        assignments=[
            ExistingPersonAssignmentInput(
                person_id=person.id,
                position="Operador",
                area="Montaje",
                workplace="Recinto nuevo",
            )
        ],
    )
    result = document_batch_service.assign_existing_people(db, payload, staff)
    assert [row["person_id"] for row in result["assigned"]] == [str(person.id)]
    assert result["already_assigned"] == []
    assert result["contracts_created"] == 1
    assert result["epp_documents_created"] == 1

    repeated = document_batch_service.assign_existing_people(db, payload, staff)
    assert repeated["assigned"] == []
    assert repeated["already_assigned"] == [{
        "person_id": str(person.id),
        "name": f"{person.first_name} {person.last_name}",
        "rut": person.rut,
    }]
    assert repeated["documents_created"] == 0
    assignments = db.scalars(
        select(DocumentEmployment).where(
            DocumentEmployment.person_id == person.id,
            DocumentEmployment.event_id == target_event.id,
        )
    ).all()
    assert len(assignments) == 1
    old_doc_after = db.get(ManagedDocument, old_doc["id"])
    assert old_doc_after is not None
    assert old_doc_after.event_id == source_event.id
    target_docs = db.scalars(
        select(ManagedDocument).where(
            ManagedDocument.person_id == person.id,
            ManagedDocument.event_id == target_event.id,
        )
    ).all()
    assert {doc.kind for doc in target_docs} == {"CONTRACT", "EPP"}


def test_event_contract_defaults_override_global_and_preflight_reports_missing_data(local_document_staff):
    db, client, staff, event = local_document_staff
    person = _person_for_staff(db, client, staff, event_id=event.id)
    global_defaults = DocumentContractDefaults(
        employer_name="Greenway Global SpA", employer_rut=_valid_rut_from_seed(87654321),
        representative_name="Representante Global", representative_rut=_valid_rut_from_seed(87654320),
        employer_address="Casa matriz 1", employer_commune="Santiago",
    )
    event_defaults = DocumentContractDefaults(
        employer_name="Greenway Faena SpA", employer_rut=_valid_rut_from_seed(76543210),
        representative_name="Representante Faena", representative_rut=_valid_rut_from_seed(76543211),
        employer_address="Oficina faena 2", employer_commune="Temuco",
    )
    managed_documents_service.save_contract_defaults(db, global_defaults, staff)
    effective = managed_documents_service.save_event_contract_defaults(db, event.id, event_defaults, staff)
    assert effective["employer_name"] == "Greenway Faena SpA"
    assert effective["has_event_override"] is True
    assert managed_documents_service.get_contract_defaults(db)["employer_name"] == "Greenway Global SpA"
    person.email = None
    db.commit()

    preview = document_batch_service.preview_missing_draft_pairs(
        db, GenerateMissingDraftsCommit(event_id=event.id, person_ids=[person.id]), staff
    )
    gaps = preview["results"][0]["missing_data"]
    assert "Correo para envío" in gaps
    assert "Fecha de inicio" in gaps
    assert "Remuneración" in gaps

    generated = document_batch_service.generate_missing_draft_pairs(
        db, GenerateMissingDraftsCommit(event_id=event.id, person_ids=[person.id]), staff
    )
    assert generated["documents_created"] == 2
    contract = db.scalar(select(ManagedDocument).where(
        ManagedDocument.person_id == person.id,
        ManagedDocument.event_id == event.id,
        ManagedDocument.kind == "CONTRACT",
    ))
    assert contract.data["employer_name"] == "Greenway Faena SpA"
    reset = managed_documents_service.clear_event_contract_defaults(db, event.id, staff)
    assert reset["employer_name"] == "Greenway Global SpA"


def test_apply_employer_defaults_only_fills_empty_fields_in_contract_drafts(local_document_staff):
    db, client, staff, event = local_document_staff
    person = _person_for_staff(db, client, staff, event_id=event.id)
    defaults = DocumentContractDefaults(
        employer_name="Greenway Servicios Ambientales SpA",
        employer_rut=_valid_rut_from_seed(76234567),
        representative_name="Gonzalo Ormeño",
        representative_rut=_valid_rut_from_seed(17188611),
        employer_address="Ñuble 1034",
        employer_commune="Santiago",
    )
    managed_documents_service.save_event_contract_defaults(db, event.id, defaults, staff)

    partial = managed_documents_service.create_document(
        db,
        DocumentCreate(
            person_id=person.id, event_id=event.id, kind="CONTRACT",
            title="Contrato parcial", data={"employer_name": "Empresa ya revisada"},
        ),
        staff,
    )
    complete = managed_documents_service.create_document(
        db,
        DocumentCreate(
            person_id=person.id, event_id=event.id, kind="CONTRACT",
            title="Contrato completo", data={
                "employer_name": "Nombre manual",
                "employer_rut": _valid_rut_from_seed(11223344),
                "representative_name": "Representante manual",
                "representative_rut": _valid_rut_from_seed(44332211),
                "employer_address": "Dirección manual",
                "employer_commune": "Temuco",
            },
        ),
        staff,
    )
    sent = managed_documents_service.create_document(
        db,
        DocumentCreate(
            person_id=person.id, event_id=event.id, kind="CONTRACT",
            title="Contrato enviado", data={},
        ),
        staff,
    )
    sent_row = db.get(ManagedDocument, sent["id"])
    sent_row.status = "SENT"
    db.commit()

    result = managed_documents_service.apply_employer_defaults_to_drafts(db, event.id, staff)
    assert result == {
        "event_id": str(event.id),
        "documents_checked": 2,
        "documents_updated": 1,
        "documents_unchanged": 1,
        "fields_filled": 5,
        "missing_defaults": [],
    }
    partial_after = db.get(ManagedDocument, partial["id"])
    assert partial_after.data["employer_name"] == "Empresa ya revisada"
    assert partial_after.data["employer_rut"] == defaults.employer_rut
    assert partial_after.data["representative_name"] == defaults.representative_name
    complete_after = db.get(ManagedDocument, complete["id"])
    assert complete_after.data["employer_name"] == "Nombre manual"
    sent_after = db.get(ManagedDocument, sent["id"])
    assert sent_after.status == "SENT"
    assert sent_after.data == {}

    repeated = managed_documents_service.apply_employer_defaults_to_drafts(db, event.id, staff)
    assert repeated["documents_updated"] == 0
    assert repeated["fields_filled"] == 0


def test_person_history_returns_actor_and_field_changes(local_document_staff):
    db, client, staff, event = local_document_staff
    person = _person_for_staff(db, client, staff, event_id=event.id)
    create_audit_log(
        db, user=staff, action="DOCUMENT_PERSON_UPDATED", module="documents",
        entity_type="DocumentPerson", entity_id=person.id, event_id=event.id,
        old_data={"first_name": "Ana"}, new_data={"first_name": "Ana María"},
    )
    history = get_person_document_history(person.id, event.id, db, staff)
    assert any(entry["action"] == "DOCUMENT_PERSON_UPDATED" for entry in history)
    entry = next(entry for entry in history if entry["action"] == "DOCUMENT_PERSON_UPDATED")
    assert entry["actor"] == staff.full_name
    assert entry["old_data"]["first_name"] == "Ana"
    assert entry["new_data"]["first_name"] == "Ana María"


def test_event_status_export_is_a_formatted_workbook_with_readable_column_widths(local_document_staff):
    db, client, staff, event = local_document_staff
    person = _person_for_staff(db, client, staff, event_id=event.id)
    content = document_reporting_service.build_event_document_status_xlsx(db, event.id, staff)
    workbook = load_workbook(BytesIO(content), read_only=False)
    sheet = workbook["Estado documental"]

    assert sheet["A1"].value == "GREENWAY APP  |  ESTADO DOCUMENTAL"
    assert sheet["A8"].value == event.name
    assert sheet["B8"].value == f"{person.first_name} {person.last_name}"
    assert sheet["H8"].value == "Faltante"
    assert sheet["J8"].value == "Faltante"
    assert sheet.freeze_panes == "A8"
    assert sheet.column_dimensions["B"].width >= 28
    assert sheet.column_dimensions["D"].width >= 30
    assert sheet.auto_filter.ref == "A7:M8"
